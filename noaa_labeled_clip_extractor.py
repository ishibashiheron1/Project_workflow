#!/usr/bin/env python3
"""
Extract annotated 5-second WAV clips + spectrogram PNGs from NOAA PIFSC DCLDE data.

Examples:
  python noaa_labeled_clip_extractor.py --excel "DCLDE2020 DetectionData.xlsx" \
      --audio-dir ./audio --dataset 1705 --kind odontocete --limit 10

  python noaa_labeled_clip_extractor.py --excel "DCLDE2020 DetectionData.xlsx" \
      --audio-dir ./audio --dataset 1705 --kind odontocete --limit 10 --cloud

The NOAA annotation is an ENCOUNTER INTERVAL, not a guarantee that the
specific 5-second sample contains a vocalization. Review before model training.
"""
import argparse
import bisect
import csv
import re
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import soundfile as sf
from scipy.signal import resample_poly
from openpyxl import load_workbook

NAME_RE = re.compile(r"(?<!\d)(1705|1706)_(\d{8})_(\d{6})_\d+\.flac$", re.I)
BUCKET = "noaa-pifsc-bioacoustic"
SHEETS = {
    ("1705", "odontocete"): "1705_OdontoceteDetections",
    ("1706", "odontocete"): "1706_OdontocetesDetections",
    ("1705", "minke"): "1705_MinkeDetections",
}

def utc(value):
    if isinstance(value, datetime):
        return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)
    raise ValueError(f"Expected an Excel datetime; got {value!r}")

def filename_start(name):
    m = NAME_RE.search(Path(name).name)
    if not m:
        return None
    return m.group(1), datetime.strptime(m.group(2) + m.group(3), "%Y%m%d%H%M%S").replace(tzinfo=timezone.utc)

def read_annotations(excel, dataset, kind):
    wb = load_workbook(excel, read_only=True, data_only=True)
    lookup = {}
    for row in wb["Species_Lookup"].iter_rows(min_row=2, values_only=True):
        if isinstance(row[0], (int, float)):
            lookup[int(row[0])] = str(row[1]).strip() if row[1] else f"species_{int(row[0])}"
    kinds = ["odontocete", "minke"] if kind == "both" else [kind]
    for category in kinds:
        sheet = SHEETS.get((dataset, category))
        if not sheet:
            continue
        rows = wb[sheet].iter_rows(values_only=True)
        headers = next(rows)
        col = {str(k): i for i, k in enumerate(headers) if k}
        for rownum, row in enumerate(rows, start=2):
            try:
                start = utc(row[col["DetectionTimeStart_UTC"]])
                end = utc(row[col["DetectionTimeEnd_UTC"]])
            except (ValueError, TypeError, IndexError):
                continue
            if end <= start:
                continue
            if category == "minke":
                species = "Minke whale"
                species_ids = "minke"
                detection_id = f"minke_row_{rownum}"
            else:
                ids = []
                for field in ("Species1ID", "Species2ID", "Species3ID"):
                    val = row[col[field]] if field in col else None
                    if isinstance(val, (int, float)):
                        ids.append(int(val))
                if not ids:
                    continue
                species_ids = "+".join(map(str, ids))
                species = " + ".join(lookup.get(i, f"Species ID {i}") for i in ids)
                detection_id = str(row[col["AcousticDetectionNumber"]])
            yield dict(start=start, end=end, category=category,
                       species=species, species_ids=species_ids,
                       detection_id=detection_id, sheet=sheet, row=rownum)

def local_index(audio_dir, dataset):
    result = {}
    if not audio_dir.exists():
        return result
    for path in audio_dir.rglob("*.flac"):
        parsed = filename_start(path.name)
        if parsed and parsed[0] == dataset:
            result[path.name] = (parsed[1], path)
    return result

def cloud_index(dataset):
    try:
        from google.cloud import storage
    except ImportError as exc:
        raise SystemExit("For --cloud, install: python -m pip install google-cloud-storage") from exc
    client = storage.Client.create_anonymous_client()
    blobs = {}
    print(f"Listing NOAA cloud objects in gs://{BUCKET}/{dataset}_FLAC/ ...", flush=True)
    for blob in client.list_blobs(BUCKET, prefix=f"{dataset}_FLAC/"):
        parsed = filename_start(blob.name)
        if parsed and parsed[0] == dataset:
            blobs[Path(blob.name).name] = (parsed[1], blob)
    return blobs

def save_outputs(source, clip_start, outbase, channel, preview_rate):
    with sf.SoundFile(source) as f:
        if channel > f.channels:
            raise ValueError(f"Channel {channel} not present ({f.channels} channels)")
        start = filename_start(Path(source).name)
        # Support downloaded local filenames with or without a folder prefix.
        if start is None:
            raise ValueError(f"Cannot infer start time from {source}")
        offset = (clip_start - start[1]).total_seconds()
        frame = round(offset * f.samplerate)
        frames = round(5 * f.samplerate)
        if frame < 0 or frame + frames > len(f):
            raise ValueError("Requested 5-second clip falls outside the FLAC recording")
        f.seek(frame)
        samples = f.read(frames, dtype="float32", always_2d=True)[:, channel - 1]
        rate = f.samplerate
    if len(samples) != frames:
        raise ValueError("Incomplete audio read")
    # Export a listening copy; preserve original high-frequency signal for spectrogram.
    if rate > preview_rate:
        from math import gcd
        d = gcd(rate, preview_rate)
        listening = resample_poly(samples, preview_rate // d, rate // d)
        out_rate = preview_rate
    else:
        listening, out_rate = samples, rate
    sf.write(str(outbase) + ".wav", listening, out_rate, subtype="PCM_16")
    fig, ax = plt.subplots(figsize=(11, 5))
    ax.specgram(samples, NFFT=4096, Fs=rate, noverlap=3072,
                cmap="magma", vmin=-125, vmax=-45)
    ax.set(xlabel="Seconds from clip start", ylabel="Frequency (kHz)",
           title=f"Channel {channel} | {clip_start:%Y-%m-%d %H:%M:%S} UTC")
    ax.set_ylim(0, min(rate / 2, 250000))
    ticks = np.linspace(0, min(rate / 2, 250000), 6)
    ax.set_yticks(ticks)
    ax.set_yticklabels([f"{v / 1000:.0f}" for v in ticks])
    fig.tight_layout()
    fig.savefig(str(outbase) + ".png", dpi=140)
    plt.close(fig)
    return rate, out_rate

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--excel", type=Path, required=True, help="DCLDE2020 DetectionData.xlsx")
    ap.add_argument("--audio-dir", type=Path, default=Path("./audio"), help="Folder of downloaded FLAC files")
    ap.add_argument("--output", type=Path, default=Path("./noaa_clips"))
    ap.add_argument("--dataset", choices=["1705", "1706"], default="1705")
    ap.add_argument("--kind", choices=["odontocete", "minke", "both"], default="odontocete")
    ap.add_argument("--limit", type=int, default=10, help="Maximum successful clips (default 10)")
    ap.add_argument("--channel", type=int, default=1, help="1-based FLAC channel number")
    ap.add_argument("--preview-rate", type=int, default=50000, help="WAV listening sample rate")
    ap.add_argument("--cloud", action="store_true", help="Download only matching files from public NOAA GCS bucket")
    args = ap.parse_args()
    if args.limit < 1 or args.channel < 1 or args.preview_rate < 1:
        ap.error("--limit, --channel, and --preview-rate must be positive")
    args.output.mkdir(parents=True, exist_ok=True)
    args.audio_dir.mkdir(parents=True, exist_ok=True)
    local = local_index(args.audio_dir, args.dataset)
    remote = cloud_index(args.dataset) if args.cloud else {}
    # Index by recording start, sorted; actual duration verified from FLAC metadata.
    entries = {}
    entries.update(remote)
    entries.update(local)
    ordered = sorted((v[0], name) for name, v in entries.items())
    starts = [s for s, _ in ordered]
    print(f"Found {len(local)} local / {len(remote)} cloud recordings for {args.dataset}")
    print(f"Output folder: {args.output.resolve()}")
    manifest = []
    skips = []
    successful = 0
    for ann in read_annotations(args.excel, args.dataset, args.kind):
        if successful >= args.limit:
            break
        # Choose the file with latest start time <= detection start.
        pos = bisect.bisect_right(starts, ann["start"]) - 1
        if pos < 0:
            skips.append((ann["sheet"], ann["row"], "No recording starts before annotation"))
            continue
        name = ordered[pos][1]
        local_path = args.audio_dir / name
        if name in local:
            local_path = local[name][1]
        elif name in remote:
            print(f"Downloading {name} (large file; may take a while)...", flush=True)
            remote[name][1].download_to_filename(str(local_path))
        else:
            skips.append((ann["sheet"], ann["row"], "Audio not downloaded; rerun with --cloud"))
            continue
        try:
            info = sf.info(local_path)
            available_end = starts[pos] + timedelta(seconds=info.duration)
            if ann["start"] + timedelta(seconds=5) > min(ann["end"], available_end):
                skips.append((ann["sheet"], ann["row"], "Less than 5 seconds inside annotation and recording"))
                continue
            label = re.sub(r"[^a-zA-Z0-9]+", "_", ann["species"]).strip("_")[:45]
            basename = f"{args.dataset}_{ann['category']}_{ann['detection_id']}_{ann['start']:%Y%m%dT%H%M%SZ}_{label}_ch{args.channel}"
            outbase = args.output / basename
            raw_rate, wav_rate = save_outputs(local_path, ann["start"], outbase, args.channel, args.preview_rate)
            manifest.append(dict(wav=basename+".wav", spectrogram=basename+".png",
                                 dataset=args.dataset, species=ann["species"],
                                 species_ids=ann["species_ids"], detection_id=ann["detection_id"],
                                 detection_start_utc=ann["start"].isoformat(),
                                 detection_end_utc=ann["end"].isoformat(),
                                 audio_file=name, channel=args.channel,
                                 original_sample_rate=raw_rate, wav_sample_rate=wav_rate,
                                 review_status="NEEDS_MANUAL_REVIEW"))
            successful += 1
            print(f"[{successful}/{args.limit}] {basename}")
        except Exception as exc:
            skips.append((ann["sheet"], ann["row"], str(exc)))
            print(f"Skipped {ann['sheet']} row {ann['row']}: {exc}", file=sys.stderr)
    if manifest:
        with (args.output / "manifest.csv").open("w", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=list(manifest[0]))
            writer.writeheader()
            writer.writerows(manifest)
    with (args.output / "skipped.csv").open("w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["sheet", "excel_row", "reason"])
        writer.writerows(skips)
    print(f"Done: {successful} WAV clips + spectrograms; {len(skips)} skipped annotations.")
    if not successful:
        print("No clips made. Confirm the matching FLAC files are in --audio-dir, or add --cloud.")

if __name__ == "__main__":
    main()
